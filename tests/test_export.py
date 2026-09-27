"""Export tests — write real files to tmp_path."""

from pathlib import Path

from pricehawk.export import export, to_csv, to_json, to_xlsx
from pricehawk.models import Product


def _products() -> list[Product]:
    return [
        Product(name="Widget", price=9.99, availability="In stock", rating="4/5",
                url="https://example.com/1", scraped_at="2026-01-01T00:00:00+00:00"),
        Product(name="Gadget", price=19.90, availability="Out of stock", rating="5/5",
                url="https://example.com/2", scraped_at="2026-01-01T00:00:00+00:00"),
    ]


def test_csv_roundtrip(tmp_path: Path):
    out = to_csv(_products(), tmp_path / "out.csv")
    lines = out.read_text(encoding="utf-8").strip().splitlines()
    assert lines[0].startswith("name,price")
    assert len(lines) == 3


def test_json_valid(tmp_path: Path):
    import json

    out = to_json(_products(), tmp_path / "out.json")
    data = json.loads(out.read_text(encoding="utf-8"))
    assert data[0]["name"] == "Widget"
    assert data[1]["price"] == 19.90


def test_xlsx_creates_sheet(tmp_path: Path):
    out = to_xlsx(_products(), tmp_path / "out.xlsx")
    assert out.exists() and out.stat().st_size > 0
    from openpyxl import load_workbook

    ws = load_workbook(out).active
    assert ws["A2"].value == "Widget"
    assert ws["B3"].value == 19.90


def test_export_dispatch(tmp_path: Path):
    assert export(_products(), tmp_path / "a.csv").suffix == ".csv"
    assert export(_products(), tmp_path / "a.json").suffix == ".json"
    assert export(_products(), tmp_path / "a.xlsx").suffix == ".xlsx"


def test_export_rejects_unknown(tmp_path: Path):
    try:
        export(_products(), tmp_path / "a.txt")
    except ValueError as exc:
        assert "unsupported" in str(exc)
    else:
        raise AssertionError("expected ValueError")
